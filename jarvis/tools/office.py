"""Office files: new Word and Excel files, and table analysis done in code.

PLAN.md skills #6 (Office & Business) and #13 (Data Analysis):
- create_docx (T1): a NEW .docx from structured blocks (python-docx).
- create_xlsx (T1): a NEW .xlsx with sheets of rows, number formats, explicit simple formulas
  and native Excel charts (openpyxl's own chart classes, no matplotlib).
- analyze_table (T0): row/column counts, per numeric column count/sum/mean/min/max and an
  optional group-by over a CSV or XLSX, computed in code with exact decimals. No model is
  ever called (hallucination control: arithmetic is done in code).

Safety:
- Every path goes through PathPolicy: allowed folders only; secret paths are T4 and refused.
- Files are created, never overwritten. The bytes are built in memory and written with an
  exclusive create, so a file that appears in the meantime is not replaced either.
- Sizes (blocks, rows, cells, text) are capped with clear errors.
- Each create is verified by re-opening the saved file and counting what is in it.
- Text cells that start with "=" stay text. Only formulas given explicitly in `formulas`
  become formulas, and only with a small allowlist of functions (no HYPERLINK, WEBSERVICE,
  INDIRECT, external links or DDE), so web-derived text can never become a live formula.

python-docx and openpyxl are imported lazily so the rest of Jarvis works without them.
"""

from __future__ import annotations

import csv
import io
import math
import re
import warnings
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from jarvis.tools.files import PathPolicy

# ---- limits ----
MAX_DOC_BLOCKS = 500
MAX_DOC_CHARS = 200_000        # all text in one document (the write_file cap)
MAX_LINE_CHARS = 300           # title and headings
MAX_LIST_ITEMS = 500
MAX_TABLE_ROWS = 1000
MAX_TABLE_COLS = 30
MAX_TABLE_CELLS = 20_000       # all tables in one document (python-docx is slow beyond this)

MAX_SHEETS = 20
MAX_SHEET_ROWS = 10_000
MAX_SHEET_COLS = 100
MAX_CELLS = 200_000            # all sheets together
MAX_CELL_CHARS = 32_767        # Excel's own limit per cell
MAX_BOOK_CHARS = 2_000_000     # all text in one workbook
MAX_FORMULAS = 500
MAX_FORMULA_CHARS = 300
MAX_CHARTS = 20
MAX_FORMAT_CHARS = 100

MAX_TABLE_BYTES = 20_000_000   # analyze_table input file
MAX_ANALYZE_ROWS = 200_000
MAX_ANALYZE_COLS = 200
MAX_GROUPS_SHOWN = 50
MAX_DISTINCT = 10_000          # distinct values counted per text column
MAX_SUMMARY_CHARS = 7000


class OfficeError(ValueError):
    """A clear refusal: bad input, a size cap, or an existing file."""


# ---- shared helpers ----

_XML_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff\ud800-\udfff]")


def _clean(text: Any) -> str:
    """Text that Word and Excel accept (control characters and lone surrogates removed)."""
    return _XML_ILLEGAL.sub("", str(text))


def _short(text: str, what: str, limit: int = MAX_LINE_CHARS) -> str:
    if len(text) > limit:
        raise OfficeError(f"{what} is too long ({len(text)} characters, limit {limit})")
    return text


def _n(count: int, word: str) -> str:
    """'1 row', '2 rows'."""
    return f"{count} {word}" if count == 1 else f"{count} {word}s"


def _cap(count: int, limit: int, what: str) -> None:
    if count > limit:
        raise OfficeError(f"too many {what}: {count} (limit {limit})")


def _new_path(policy: PathPolicy, raw: str, suffix: str) -> Path:
    path = policy.resolve(raw)
    if path.suffix.lower() != suffix:
        raise OfficeError(f"{raw} must end with {suffix}")
    if path.exists():
        raise OfficeError(f"{raw} already exists; office tools only create new files (choose another name)")
    return path


def _write_new(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with open(path, "xb") as f:  # exclusive create: never replaces a file
            f.write(data)
    except FileExistsError:
        raise OfficeError(f"{path.name} already exists; office tools only create new files") from None


def _import_docx():
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover - dependency missing
        raise OfficeError("python-docx is not installed (pip install -e .)") from e
    return docx


def _import_openpyxl():
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover - dependency missing
        raise OfficeError("openpyxl is not installed (pip install -e .)") from e
    return openpyxl


def _plain(value: Any) -> str:
    """A cell value as display text (integral floats without ".0", dates as ISO)."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return _clean(value)


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def _col_letter(n: int) -> str:
    out = ""
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out


_CELL_RE = re.compile(r"\$?([A-Z]{1,3})\$?([1-9][0-9]{0,6})")
_RANGE_RE = re.compile(r"\$?([A-Z]{1,3})\$?([1-9][0-9]{0,6}):\$?([A-Z]{1,3})\$?([1-9][0-9]{0,6})")


def _cell_ref(raw: str, what: str) -> tuple[int, int]:
    """'B7' -> (row 7, column 2)."""
    m = _CELL_RE.fullmatch(str(raw).strip().upper())
    if not m or _col_index(m.group(1)) > 16384:
        raise OfficeError(f"{what} {raw!r} is not a cell like B7")
    return int(m.group(2)), _col_index(m.group(1))


# ---- create_docx ----

DOCX_PARAMS = {
    "type": "object",
    "properties": {
        "path": {"type": "string", "description": "New file name ending in .docx"},
        "title": {"type": "string", "description": "Document title (optional)"},
        "blocks": {"type": "array", "description": "Content in order", "items": {
            "type": "object",
            "properties": {
                "type": {"type": "string", "enum": ["heading", "paragraph", "bullets", "numbered", "table"]},
                "text": {"type": "string", "description": "heading or paragraph text"},
                "level": {"type": "integer", "description": "heading level 1-3"},
                "items": {"type": "array", "items": {"type": "string"}, "description": "list items"},
                "header": {"type": "array", "items": {"type": "string"}, "description": "table header cells"},
                "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}},
                         "description": "table rows, each a list of cells"},
            },
            "required": ["type"]}},
    },
    "required": ["path", "blocks"],
}


@dataclass
class _DocPlan:
    """Normalized document content plus the counts the verifier expects."""
    title: str = ""
    blocks: list[dict] = field(default_factory=list)
    headings: int = 0
    paragraphs: int = 0
    bullets: int = 0
    numbered: int = 0
    tables: list[int] = field(default_factory=list)  # rows per table, header included

    def describe(self) -> str:
        parts = (["a title"] if self.title else []) + [
            _n(n, what) for n, what in ((self.headings, "heading"), (self.paragraphs, "paragraph"),
                                        (self.bullets + self.numbered, "list item"),
                                        (len(self.tables), "table")) if n]
        return ", ".join(parts)


def _doc_plan(args: dict) -> _DocPlan:
    plan = _DocPlan(title=_short(_clean(args.get("title") or "").strip(), "the title"))
    blocks = args.get("blocks") or []
    _cap(len(blocks), MAX_DOC_BLOCKS, "blocks")
    if not plan.title and not blocks:
        raise OfficeError("nothing to write: give a title or at least one block")
    total, cells = len(plan.title), 0
    for i, b in enumerate(blocks, 1):
        kind = b.get("type")
        where = f"block {i} ({kind})"
        if kind in ("heading", "paragraph"):
            text = _clean(b.get("text") or "").strip()
            if not text:
                raise OfficeError(f"{where} has no text")
            if kind == "heading":
                level = b.get("level", 1)
                if level not in (1, 2, 3) or isinstance(level, bool):
                    raise OfficeError(f"{where}: the level must be 1, 2 or 3")
                plan.blocks.append({"type": kind, "text": _short(text, where), "level": int(level)})
                plan.headings += 1
            else:
                plan.blocks.append({"type": kind, "text": text})
                plan.paragraphs += 1
            total += len(text)
        elif kind in ("bullets", "numbered"):
            items = [t for t in (_clean(x).strip() for x in b.get("items") or []) if t]
            if not items:
                raise OfficeError(f"{where} has no items")
            _cap(len(items), MAX_LIST_ITEMS, f"items in {where}")
            plan.blocks.append({"type": kind, "items": items})
            if kind == "bullets":
                plan.bullets += len(items)
            else:
                plan.numbered += len(items)
            total += sum(map(len, items))
        elif kind == "table":
            header = [_plain(x).strip() for x in b.get("header") or []]
            if not header:
                raise OfficeError(f"{where} needs a header row")
            _cap(len(header), MAX_TABLE_COLS, f"columns in {where}")
            rows_in = b.get("rows") or []
            _cap(len(rows_in), MAX_TABLE_ROWS, f"rows in {where}")
            rows = []
            for r, row in enumerate(rows_in, 1):
                if len(row) > len(header):
                    raise OfficeError(f"{where}: row {r} has {len(row)} cells but the header has {len(header)}")
                rows.append([_plain(x) for x in row] + [""] * (len(header) - len(row)))
            cells += len(header) * (1 + len(rows))
            _cap(cells, MAX_TABLE_CELLS, "table cells in the document")
            total += sum(len(c) for row in (header, *rows) for c in row)
            plan.blocks.append({"type": kind, "header": header, "rows": rows})
            plan.tables.append(1 + len(rows))
        else:
            raise OfficeError(f"{where}: unknown block type")
        if total > MAX_DOC_CHARS:
            raise OfficeError(f"the document text is too long (over {MAX_DOC_CHARS} characters)")
    return plan


def _fresh_numbering(doc) -> int | None:
    """A new numbering instance of the "List Number" style that restarts at 1, so a second
    numbered list does not continue the first one's numbers. None if the template lacks it."""
    try:
        num_id = doc.styles["List Number"].element.pPr.numPr.numId.val
        numbering = doc.part.numbering_part.element
        num = numbering.add_num(numbering.num_having_numId(num_id).abstractNumId.val)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId
    except (AttributeError, KeyError):
        return None


def _set_numbering(paragraph, num_id: int) -> None:
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


# ---- create_xlsx ----

XLSX_PARAMS = {
    "type": "object",
    "properties": {
        "path": {"type": "string", "description": "New file name ending in .xlsx"},
        "sheets": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "name": {"type": "string"},
                "rows": {"type": "array", "items": {"type": "array"},
                         "description": "First row = header. Cells: text, numbers, true/false or null"},
            },
            "required": ["rows"]}},
        "formats": {"type": "array", "description": "Number formats for data cells, e.g. column B '#,##0.00'",
                    "items": {"type": "object", "properties": {
                        "sheet": {"type": "string"},
                        "column": {"type": "string", "description": "column letter or header name"},
                        "format": {"type": "string"}}, "required": ["column", "format"]}},
        "formulas": {"type": "array", "description": "Only when asked for: e.g. cell B8, formula =SUM(B2:B7)",
                     "items": {"type": "object", "properties": {
                         "sheet": {"type": "string"}, "cell": {"type": "string"}, "formula": {"type": "string"}},
                         "required": ["cell", "formula"]}},
        "charts": {"type": "array", "items": {"type": "object", "properties": {
            "sheet": {"type": "string"},
            "type": {"type": "string", "enum": ["bar", "line", "pie"]},
            "title": {"type": "string"},
            "data": {"type": "string", "description": "value columns with their header row, e.g. B1:C7"},
            "categories": {"type": "string", "description": "label column without header, e.g. A2:A7"},
            "anchor": {"type": "string", "description": "top-left cell of the chart, e.g. E2"}},
            "required": ["type", "data"]}},
    },
    "required": ["path", "sheets"],
}

# Formulas: classic functions only (no prefix needed in the file), nothing that reaches the
# network, other files or other programs.
SAFE_FUNCTIONS = frozenset({"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "COUNTBLANK", "COUNTIF", "SUMIF",
                            "AVERAGEIF", "MEDIAN", "PRODUCT", "ROUND", "ROUNDUP", "ROUNDDOWN", "ABS", "IF",
                            "IFERROR", "AND", "OR", "NOT"})
_FORMULA_STRING = re.compile(r'"[^"\r\n]{0,100}"')
_FORMULA_TOKEN = re.compile(
    r"(?P<space>\s+)"
    r"|(?P<func>[A-Za-z][A-Za-z0-9]*)\("
    r"|(?P<ref>(?:(?P<sheet>[A-Za-z_][A-Za-z0-9_]{0,30})!)?"
    r"\$?[A-Za-z]{1,3}\$?[0-9]{1,7}(?::\$?[A-Za-z]{1,3}\$?[0-9]{1,7})?)(?![A-Za-z0-9_(!])"
    r"|(?P<num>[0-9]+(?:\.[0-9]+)?%?)"
    r"|(?P<bool>TRUE|FALSE)(?![A-Za-z0-9_(])"
    r"|(?P<op><>|<=|>=|[-+*/^&(),<>=])")
_BAD_SHEET_CHARS = set(":\\/?*[]")


def check_formula(raw: str, sheet_names: set[str]) -> str:
    """The formula with "=" in front, function names and cell references in capitals, or
    OfficeError if it uses anything outside the simple allowlist."""
    text = str(raw).strip()
    text = text[1:].strip() if text.startswith("=") else text
    if not text:
        raise OfficeError("empty formula")
    _short(text, "the formula", MAX_FORMULA_CHARS)
    # String literals are kept as they are; blank them out (same length) for the token scan.
    body = _FORMULA_STRING.sub(lambda m: "0" * len(m.group()), text)
    allowed = ", ".join(sorted(SAFE_FUNCTIONS))
    malformed = OfficeError(f"formula ={text} is not well formed")
    out, pos, depth = [], 0, 0
    prev = "start"  # start | operand | operator | open | comma: a minimal grammar, so Excel never
    while pos < len(body):  # has to "repair" the file because of a malformed formula
        m = _FORMULA_TOKEN.match(body, pos)
        if not m:
            raise OfficeError(f"formula ={text} is not allowed near {text[pos:pos + 12]!r}: only cell references, "
                              f"numbers, quoted text, + - * / ^ & comparisons and these functions: {allowed}")
        piece, pos = text[m.start():m.end()], m.end()
        op = m.group("op")
        if m.group("space"):
            continue
        if m.group("func") or op == "(":
            if prev == "operand":
                raise malformed
            if m.group("func"):
                name = m.group("func").upper()
                if name not in SAFE_FUNCTIONS:
                    raise OfficeError(f"function {name} is not allowed in formulas (allowed: {allowed})")
                piece = name + "("
            depth, prev = depth + 1, "open"
        elif op == ")":
            depth -= 1
            if depth < 0 or prev in ("operator", "comma"):
                raise malformed
            prev = "operand"
        elif op == ",":
            if depth == 0 or prev != "operand":
                raise malformed
            prev = "comma"
        elif op:
            unary = op in "+-" and prev != "operand"
            if prev != "operand" and not unary:
                raise malformed
            prev = "operator"
        else:  # a reference, number, quoted text or TRUE/FALSE
            if prev == "operand":
                raise malformed
            if m.group("ref"):
                sheet = m.group("sheet")
                if sheet is not None and sheet.lower() not in sheet_names:
                    raise OfficeError(f"formula ={text} refers to an unknown sheet {sheet!r}")
                prefix = f"{sheet}!" if sheet is not None else ""
                piece = prefix + piece[len(prefix):].upper()
            prev = "operand"
        out.append(piece)
    if depth or prev != "operand":
        raise malformed
    return "=" + "".join(out)


def _xl_value(v: Any, where: str) -> Any:
    if v is None or isinstance(v, bool) or isinstance(v, int):
        return v
    if isinstance(v, float):
        if not math.isfinite(v):
            raise OfficeError(f"{where}: numbers must be finite")
        return v
    if isinstance(v, str):
        return _short(_clean(v), where, MAX_CELL_CHARS)
    raise OfficeError(f"{where}: cells must be text, numbers, true/false or null")


def _filled(v: Any) -> bool:
    return v is not None and v != ""


@dataclass
class _Sheet:
    name: str
    rows: list[list[Any]]
    formulas: dict[tuple[int, int], str] = field(default_factory=dict)  # (row, col) -> formula
    charts: int = 0

    @property
    def width(self) -> int:
        return max([len(r) for r in self.rows] + [c for _, c in self.formulas])

    def filled_rows(self) -> list[int]:
        rows = {i for i, row in enumerate(self.rows, 1) if any(_filled(v) for v in row)}
        return sorted(rows | {r for r, _ in self.formulas})

    @property
    def last_row(self) -> int:
        filled = self.filled_rows()
        return filled[-1] if filled else 1

    def value_at(self, row: int, col: int) -> Any:
        if row <= len(self.rows) and col <= len(self.rows[row - 1]):
            return self.rows[row - 1][col - 1]
        return None

    def column(self, raw: str) -> int:
        """A column by header name (case-insensitive) or by letter."""
        want = " ".join(str(raw).split()).lower()
        for i, v in enumerate(self.rows[0], 1):
            if _filled(v) and " ".join(_plain(v).split()).lower() == want:
                return i
        if re.fullmatch(r"[A-Za-z]{1,3}", want) and _col_index(want.upper()) <= self.width:
            return _col_index(want.upper())
        raise OfficeError(f"sheet {self.name!r} has no column {raw!r}")

    def range(self, raw: str, what: str) -> tuple[int, int, int, int]:
        """'B1:C7' -> (min_col, min_row, max_col, max_row), inside the filled area."""
        m = _RANGE_RE.fullmatch(str(raw).strip().upper())
        if not m:
            raise OfficeError(f"{what} {raw!r} is not a range like B1:C7")
        c1, c2 = sorted((_col_index(m.group(1)), _col_index(m.group(3))))
        r1, r2 = sorted((int(m.group(2)), int(m.group(4))))
        if c2 > self.width or r2 > self.last_row:
            raise OfficeError(f"{what} {raw} is outside the data of sheet {self.name!r} "
                              f"(A1:{_col_letter(self.width)}{self.last_row})")
        return c1, r1, c2, r2


@dataclass
class _Chart:
    sheet: _Sheet
    kind: str
    title: str
    data: tuple[int, int, int, int]
    categories: tuple[int, int, int, int] | None
    anchor: str


@dataclass
class _Book:
    sheets: list[_Sheet]
    formats: list[tuple[_Sheet, int, str]] = field(default_factory=list)
    charts: list[_Chart] = field(default_factory=list)

    def sheet(self, raw: str | None) -> _Sheet:
        if not raw:
            return self.sheets[0]
        for s in self.sheets:
            if s.name.lower() == str(raw).strip().lower():
                return s
        raise OfficeError(f"no sheet named {raw!r} (sheets: {', '.join(s.name for s in self.sheets)})")

    def describe(self) -> str:
        parts = [", ".join(f"sheet {s.name} ({_n(len(s.filled_rows()), 'row')})" for s in self.sheets)]
        formulas = sum(len(s.formulas) for s in self.sheets)
        if formulas:
            parts.append(_n(formulas, "formula"))
        if self.charts:
            kinds = Counter(c.kind for c in self.charts)
            parts.append(f"{_n(len(self.charts), 'chart')} ({', '.join(f'{n} {k}' for k, n in sorted(kinds.items()))})")
        return "; ".join(parts)


def _sheet_name(raw: Any, i: int, seen: set[str]) -> str:
    name = _clean(raw if raw else f"Sheet{i}").strip()
    if not name or len(name) > 31 or any(ch in _BAD_SHEET_CHARS for ch in name) or name[0] == "'" or name[-1] == "'":
        raise OfficeError(f"sheet name {name!r} is not allowed in Excel "
                          "(1 to 31 characters, none of : \\ / ? * [ ], no quote at the start or end)")
    if name.lower() == "history":
        raise OfficeError("Excel reserves the sheet name 'History'")
    if name.lower() in seen:
        raise OfficeError(f"two sheets are named {name!r}")
    seen.add(name.lower())
    return name


def _book_plan(args: dict) -> _Book:
    sheets_in = args.get("sheets") or []
    if not sheets_in:
        raise OfficeError("give at least one sheet")
    _cap(len(sheets_in), MAX_SHEETS, "sheets")
    seen: set[str] = set()
    cells = chars = 0
    book = _Book([])
    for i, s in enumerate(sheets_in, 1):
        name = _sheet_name(s.get("name"), i, seen)
        rows_in = s.get("rows") or []
        if not rows_in:
            raise OfficeError(f"sheet {name!r} has no rows (the first row is the header)")
        _cap(len(rows_in), MAX_SHEET_ROWS, f"rows in sheet {name!r}")
        rows = []
        for r, row in enumerate(rows_in, 1):
            if not isinstance(row, list):
                raise OfficeError(f"sheet {name!r} row {r} must be a list of cells")
            _cap(len(row), MAX_SHEET_COLS, f"cells in row {r} of sheet {name!r}")
            rows.append([_xl_value(v, f"sheet {name!r} row {r}") for v in row])
            cells += len(row)
            chars += sum(len(v) for v in rows[-1] if isinstance(v, str))
        _cap(cells, MAX_CELLS, "cells in the workbook")
        _cap(chars, MAX_BOOK_CHARS, "characters of text in the workbook")
        if not any(_filled(v) for v in rows[0]):
            raise OfficeError(f"sheet {name!r}: the first row (the header) is empty")
        book.sheets.append(_Sheet(name, rows))

    formulas = args.get("formulas") or []
    _cap(len(formulas), MAX_FORMULAS, "formulas")
    for f in formulas:
        sheet = book.sheet(f.get("sheet"))
        row, col = _cell_ref(f["cell"], "formula cell")
        if row > MAX_SHEET_ROWS + 100 or col > MAX_SHEET_COLS:
            raise OfficeError(f"formula cell {f['cell']} is too far from the data")
        if _filled(sheet.value_at(row, col)) or (row, col) in sheet.formulas:
            raise OfficeError(f"cell {sheet.name}!{f['cell']} already holds a value; put the formula in an empty cell")
        sheet.formulas[(row, col)] = check_formula(f["formula"], seen)

    for f in args.get("formats") or []:
        sheet = book.sheet(f.get("sheet"))
        fmt = str(f["format"])
        if not fmt.strip() or "\n" in fmt or "\r" in fmt or fmt.count('"') % 2:
            raise OfficeError(f"number format {fmt!r} is not valid")
        book.formats.append((sheet, sheet.column(f["column"]), _short(fmt, "the number format", MAX_FORMAT_CHARS)))

    charts = args.get("charts") or []
    _cap(len(charts), MAX_CHARTS, "charts")
    for c in charts:
        sheet = book.sheet(c.get("sheet"))
        kind = c["type"]
        data = sheet.range(c["data"], "chart data")
        if data[3] - data[1] < 1:
            raise OfficeError(f"chart data {c['data']} needs its header row and at least one row of values")
        if kind == "pie" and data[0] != data[2]:
            raise OfficeError("a pie chart takes one column of values")
        cats = sheet.range(c["categories"], "chart categories") if c.get("categories") else None
        if cats and cats[0] != cats[2]:
            raise OfficeError("chart categories must be one column")
        if c.get("anchor"):
            row, col = _cell_ref(c["anchor"], "chart anchor")
            anchor = f"{_col_letter(col)}{row}"
        else:  # right of the data, charts on one sheet stacked downwards
            anchor = f"{_col_letter(sheet.width + 2)}{2 + 18 * sheet.charts}"
        sheet.charts += 1
        title = _short(_clean(c.get("title") or "").strip(), "the chart title")
        book.charts.append(_Chart(sheet, kind, title, data, cats, anchor))
    return book


_CHART_TAG = re.compile(r"<(?:\w+:)?(barChart|lineChart|pieChart)\b")
_CHART_KIND = {"barChart": "bar", "lineChart": "line", "pieChart": "pie"}


def chart_kinds(path: Path) -> list[str]:
    """Chart types stored in an .xlsx file, read from its chart parts."""
    with zipfile.ZipFile(path) as z:
        kinds = []
        for name in z.namelist():
            if re.fullmatch(r"xl/charts/chart\d+\.xml", name):
                m = _CHART_TAG.search(z.read(name).decode("utf-8", errors="replace"))
                kinds.append(_CHART_KIND[m.group(1)] if m else "other")
        return sorted(kinds)


# ---- analyze_table ----

ANALYZE_PARAMS = {
    "type": "object",
    "properties": {
        "path": {"type": "string", "description": "A .csv or .xlsx file"},
        "sheet": {"type": "string", "description": "XLSX sheet name (default: the first sheet)"},
        "group_by": {"type": "string", "description": "column to group rows by (optional)"},
        "sum_column": {"type": "string", "description": "numeric column to sum per group (optional)"},
    },
    "required": ["path"],
}

_BLANKS = {"", "-", "--", "—", "–", "na", "n/a", "nil", "null", "none", "nan"}
_CURRENCY = r"(?:₹|rs\.?|inr|\$|€|£)"
_NUMBER_RE = re.compile(
    rf"(?P<sign1>[-+])?\s*(?:{_CURRENCY}\s*)?(?P<sign2>[-+])?\s*"
    rf"(?P<int>[0-9][0-9,]*)?(?P<frac>\.[0-9]+)?\s*(?:{_CURRENCY})?\s*(?:/-)?", re.I)
_GROUPED_WESTERN = re.compile(r"[0-9]{1,3}(?:,[0-9]{3})+")              # 1,234,567
_GROUPED_INDIAN = re.compile(r"[0-9]{1,2}(?:,[0-9]{2})*,[0-9]{3}")      # 12,34,567


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip().lower() in _BLANKS)


def parse_number(text: str) -> Decimal | None:
    """A number written the way people type it in India, or None.

    Accepts plain numbers, Indian grouping ("1,23,456.50"), western grouping ("123,456.50"),
    a currency mark before or after (₹, Rs, Rs., INR, $, €, £), a trailing "/-", a sign
    (also the Unicode minus) and accounting brackets for negatives ("(1,200)"). Commas in
    any other pattern ("1,2,3") are not a number."""
    s = " ".join(str(text).replace("−", "-").split())
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1].strip()
    m = _NUMBER_RE.fullmatch(s)
    if not m or not (m.group("int") or m.group("frac")):
        return None
    signs = [g for g in (m.group("sign1"), m.group("sign2")) if g]
    if len(signs) > 1 or (signs and negative):
        return None
    whole = m.group("int") or "0"
    if "," in whole and not (_GROUPED_WESTERN.fullmatch(whole) or _GROUPED_INDIAN.fullmatch(whole)):
        return None
    try:
        number = Decimal(whole.replace(",", "") + (m.group("frac") or ""))
    except InvalidOperation:  # pragma: no cover - the pattern only lets digits through
        return None
    return -number if negative or signs == ["-"] else number


def to_number(value: Any) -> Decimal | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(repr(value)) if math.isfinite(value) else None
    if isinstance(value, str):
        return parse_number(value)
    return None


def fmt_number(d: Decimal) -> str:
    """Plain digits, at most 4 decimals, no exponent."""
    try:
        if d == d.to_integral_value():
            return format(d.quantize(Decimal(1)), "f")
        text = format(d.quantize(Decimal("0.0001"), ROUND_HALF_UP), "f")
        return text.rstrip("0").rstrip(".")
    except InvalidOperation:  # more digits than the decimal context holds
        return str(d)


def _snippet(value: Any, limit: int = 40) -> str:
    text = " ".join(_plain(value).split())
    return text if len(text) <= limit else text[:limit - 1] + "…"


@dataclass
class _Column:
    name: str
    numbers: list[Decimal] = field(default_factory=list)
    texts: int = 0
    dates: int = 0
    blanks: int = 0
    distinct: set[str] = field(default_factory=set)

    @property
    def kind(self) -> str:
        filled = len(self.numbers) + self.texts + self.dates
        if not filled:
            return "empty"
        if self.numbers and len(self.numbers) * 2 >= filled:
            return "number"
        return "date" if self.dates * 2 >= filled else "text"


def _read_csv(path: Path) -> tuple[list[list[Any]], str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost anything
        text = raw.decode("utf-8", errors="replace")
    first = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max(",;\t|", key=first.count) if any(d in first for d in ",;\t|") else ","
    rows: list[list[Any]] = []
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):
            _cap(len(rows), MAX_ANALYZE_ROWS, "rows to analyze")
            _cap(len(row), MAX_ANALYZE_COLS, "columns to analyze")
            rows.append(row)
    except csv.Error as e:
        raise OfficeError(f"cannot read {path.name} as CSV: {e}") from None
    names = {",": "comma", ";": "semicolon", "\t": "tab", "|": "pipe"}
    return rows, f"CSV, {names[delimiter]}-separated"


def _read_xlsx(path: Path, sheet: str | None) -> tuple[list[list[Any]], str]:
    openpyxl = _import_openpyxl()
    with warnings.catch_warnings():  # unsupported extensions in other people's files only warn
        warnings.simplefilter("ignore")
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # openpyxl raises many types for broken files
            raise OfficeError(f"cannot read {path.name} as an Excel workbook: {e}") from None
        try:
            sheets = wb.worksheets
            if not sheets:
                raise OfficeError(f"{path.name} has no worksheets")
            names = [ws.title for ws in sheets]
            if sheet:
                found = [ws for ws in sheets if ws.title.lower() == sheet.strip().lower()]
                if not found:
                    raise OfficeError(f"no sheet named {sheet!r} (sheets: {', '.join(names)})")
                ws = found[0]
            else:
                ws = sheets[0]
            # data_only: the values Excel saved for formula cells. Files written by openpyxl
            # (create_xlsx included) have no saved values, so their formula cells read as blank.
            rows: list[list[Any]] = []
            for row in ws.iter_rows(values_only=True):
                _cap(len(rows), MAX_ANALYZE_ROWS, "rows to analyze")
                values = list(row)
                while values and values[-1] is None:
                    values.pop()
                _cap(len(values), MAX_ANALYZE_COLS, "columns to analyze")
                rows.append(values)
        finally:
            wb.close()
    where = f"XLSX, sheet {ws.title}" + (f" ({names.index(ws.title) + 1} of {len(names)}: {', '.join(names)})"
                                         if len(names) > 1 else "")
    return rows, where


_TOTAL_LABEL = re.compile(r"(grand |sub ?)?total\b", re.I)


def analyze_rows(rows: list[list[Any]], group_by: str | None = None, sum_column: str | None = None,
                 label: str = "table") -> str:
    """The text summary for a table whose first non-blank row is the header."""
    if sum_column and not group_by:
        raise OfficeError("sum_column needs group_by (the column to group rows by)")
    numbered = [(n, r) for n, r in enumerate(rows, 1) if any(not is_blank(v) for v in r)]
    skipped = len(rows) - len(numbered)
    if not numbered:
        return f"Table: {label}\nThe table is empty."
    header, data = numbered[0][1], [r for _, r in numbered[1:]]
    width = max(len(r) for _, r in numbered)
    names: list[str] = []
    for i in range(width):
        base = _snippet(header[i], 60) if i < len(header) and not is_blank(header[i]) else f"column {i + 1}"
        name, n = base, 2
        while name.lower() in (x.lower() for x in names):
            name, n = f"{base} ({n})", n + 1
        names.append(name)
    cols = [_Column(n) for n in names]
    totals: list[tuple[int, str]] = []
    for line_no, row in numbered[1:]:
        for i, col in enumerate(cols):
            v = row[i] if i < len(row) else None
            if is_blank(v):
                col.blanks += 1
                continue
            number = None if isinstance(v, (datetime, date, time)) else to_number(v)
            if number is not None:
                col.numbers.append(number)
            elif isinstance(v, (datetime, date, time)):
                col.dates += 1
            else:
                col.texts += 1
            if len(col.distinct) <= MAX_DISTINCT:
                col.distinct.add(" ".join(_plain(v).split()))
        first = next((v for v in row if not is_blank(v)), None)
        if isinstance(first, str) and _TOTAL_LABEL.match(first.strip()):
            totals.append((line_no, _snippet(first)))

    def find(raw: str, what: str) -> int:
        want = " ".join(raw.split()).lower()
        for i, n in enumerate(names):
            if n.lower() == want:
                return i
        if re.fullmatch(r"[a-z]{1,2}", want) and _col_index(want.upper()) <= width:
            return _col_index(want.upper()) - 1
        raise OfficeError(f"no {what} column {raw!r} (columns: {', '.join(names)})")

    out = [f"Table: {label}",
           f"Rows: {_n(len(data), 'data row')} (header excluded), {_n(width, 'column')}."
           + (f" {_n(skipped, 'blank row')} skipped." if skipped else ""),
           "Columns:"]
    for col in cols:
        kind = col.kind
        line = f"- {col.name}: {kind}"
        if kind == "number":
            nums = col.numbers
            total = sum(nums, Decimal(0))
            line += f", {_n(len(nums), 'value')}"
            if col.texts or col.dates:
                line += f", {col.texts + col.dates} non-numeric skipped"
            line += (f"; sum {fmt_number(total)}, mean {fmt_number(total / len(nums))}, "
                     f"min {fmt_number(min(nums))}, max {fmt_number(max(nums))}")
        elif kind != "empty":
            distinct = f"over {MAX_DISTINCT}" if len(col.distinct) > MAX_DISTINCT else str(len(col.distinct))
            line += f", {_n(len(col.numbers) + col.texts + col.dates, 'value')}, {distinct} distinct"
        if col.blanks and kind != "empty":
            line += f", {col.blanks} blank"
        out.append(line)

    if group_by:
        g = find(group_by, "group_by")
        v = find(sum_column, "sum_column") if sum_column else None
        if v is not None and cols[v].kind != "number":
            raise OfficeError(f"column {names[v]!r} is not numeric, so it cannot be summed")
        groups: dict[str, list] = {}
        for row in data:
            key_raw = row[g] if g < len(row) else None
            key = "(blank)" if is_blank(key_raw) else " ".join(_plain(key_raw).split())  # never truncated
            entry = groups.setdefault(key, [Decimal(0), 0, 0])  # sum, rows, rows with a number
            entry[1] += 1
            if v is not None:
                number = to_number(row[v]) if v < len(row) else None
                if number is not None:
                    entry[0] += number
                    entry[2] += 1
        order = sorted(groups.items(), key=lambda kv: (-kv[1][0], -kv[1][1], kv[0]) if v is not None
                       else (-kv[1][1], kv[0]))
        out.append(f"Group by {names[g]}" + (f" (sum of {names[v]}):" if v is not None else " (row count):"))
        for key, (total, count, with_value) in order[:MAX_GROUPS_SHOWN]:
            if v is None:
                out.append(f"- {_snippet(key)}: {_n(count, 'row')}")
            else:
                note = f" ({with_value} with a value)" if with_value != count else ""
                out.append(f"- {_snippet(key)}: sum {fmt_number(total)}, {_n(count, 'row')}{note}")
        rest = order[MAX_GROUPS_SHOWN:]
        if rest:
            more = f"... and {len(rest)} more groups: {_n(sum(e[1] for _, e in rest), 'row')}"
            if v is not None:
                more += f", sum {fmt_number(sum((e[0] for _, e in rest), Decimal(0)))}"
            out.append(more)
        all_rows = sum(e[1] for e in groups.values())
        if v is not None:
            out.append(f"Total: sum {fmt_number(sum((e[0] for e in groups.values()), Decimal(0)))} "
                       f"over {_n(all_rows, 'row')} in {_n(len(groups), 'group')}")
        else:
            out.append(f"Total: {_n(all_rows, 'row')} in {_n(len(groups), 'group')}")
    for line_no, text in totals[:5]:
        out.append(f"Note: row {line_no} looks like a total row ({text!r}); it is included in the sums above.")
    out.append("All numbers above were computed in code from the file.")
    text = "\n".join(out)
    if len(text) > MAX_SUMMARY_CHARS:
        text = text[:MAX_SUMMARY_CHARS] + "\n[... summary truncated]"
    return text


# ---- the tools ----

class OfficeTools:
    def __init__(self, policy: PathPolicy):
        self.policy = policy

    # ---- T1: create only ----

    def create_docx(self, args: dict) -> str:
        path = _new_path(self.policy, args["path"], ".docx")
        plan = _doc_plan(args)
        docx = _import_docx()
        doc = docx.Document()
        doc.core_properties.author = "Jarvis"
        if plan.title:
            doc.core_properties.title = plan.title
            doc.add_heading(plan.title, level=0)
        for b in plan.blocks:
            kind = b["type"]
            if kind == "heading":
                doc.add_heading(b["text"], level=b["level"])
            elif kind == "paragraph":
                doc.add_paragraph(b["text"])
            elif kind == "bullets":
                for item in b["items"]:
                    doc.add_paragraph(item, style="List Bullet")
            elif kind == "numbered":
                num_id = _fresh_numbering(doc)
                for item in b["items"]:
                    p = doc.add_paragraph(item, style="List Number")
                    if num_id is not None:
                        _set_numbering(p, num_id)
            else:
                table = doc.add_table(rows=1, cols=len(b["header"]))
                table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, b["header"]):
                    cell.paragraphs[0].add_run(text).bold = True
                for row in b["rows"]:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = text
        buf = io.BytesIO()
        doc.save(buf)
        _write_new(path, buf.getvalue())
        return f"Created {self.policy.show(path)}: {plan.describe()}."

    def verify_docx(self, args: dict) -> tuple[bool, str]:
        path = self.policy.resolve(args["path"])
        plan = _doc_plan(args)
        docx = _import_docx()
        try:
            doc = docx.Document(str(path))
        except Exception as e:  # python-docx raises several types for broken files
            return False, f"the file does not open as a Word document: {e}"
        styles = Counter((p.style.name if p.style is not None else "") for p in doc.paragraphs)
        got = {"title": styles["Title"], "headings": sum(styles[f"Heading {n}"] for n in (1, 2, 3)),
               "tables": [len(t.rows) for t in doc.tables], "bullet items": styles["List Bullet"],
               "numbered items": styles["List Number"]}
        want = {"title": int(bool(plan.title)), "headings": plan.headings, "tables": plan.tables,
                "bullet items": plan.bullets, "numbered items": plan.numbered}
        diffs = [f"{k} {got[k]} (expected {want[k]})" for k in want if got[k] != want[k]]
        if diffs:
            return False, "the document opens but " + ", ".join(diffs)
        return True, (f"the document opens with {_n(plan.headings, 'heading')}, {_n(len(plan.tables), 'table')} "
                      f"and {_n(plan.bullets + plan.numbered, 'list item')} as expected")

    def create_xlsx(self, args: dict) -> str:
        path = _new_path(self.policy, args["path"], ".xlsx")
        book = _book_plan(args)
        openpyxl = _import_openpyxl()
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        from openpyxl.styles import Font

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        bold = Font(bold=True)
        for s in book.sheets:
            ws = wb.create_sheet(s.name)
            widths: dict[int, int] = {}
            for r, row in enumerate(s.rows, 1):
                for c, value in enumerate(row, 1):
                    if value is None:
                        continue
                    cell = ws.cell(row=r, column=c, value=value)
                    if cell.data_type == "f":
                        cell.data_type = "s"  # text that starts with "=" stays text, never a formula
                    if r == 1:
                        cell.font = bold
                    if r <= 500:
                        widths[c] = max(widths.get(c, 0), len(_plain(value)))
            for (r, c), formula in s.formulas.items():
                ws.cell(row=r, column=c, value=formula)
            for c, w in widths.items():
                ws.column_dimensions[_col_letter(c)].width = min(max(8, w + 2), 60)
            ws.freeze_panes = "A2"
        for s, col, fmt in book.formats:
            ws = wb[s.name]
            for r in range(2, s.last_row + 1):
                ws.cell(row=r, column=col).number_format = fmt
        for ch in book.charts:
            ws = wb[ch.sheet.name]
            chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[ch.kind]()
            if ch.kind == "bar":
                chart.type = "col"
            if ch.kind != "pie":  # openpyxl 3.1 leaves axes hidden in newer Excel unless set
                chart.x_axis.delete = False
                chart.y_axis.delete = False
            if ch.title:
                chart.title = ch.title
            c1, r1, c2, r2 = ch.data
            chart.add_data(Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2), titles_from_data=True)
            if ch.categories:
                c1, r1, c2, r2 = ch.categories
                chart.set_categories(Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2))
            ws.add_chart(chart, ch.anchor)
        buf = io.BytesIO()
        wb.save(buf)
        _write_new(path, buf.getvalue())
        return f"Created {self.policy.show(path)}: {book.describe()}."

    def verify_xlsx(self, args: dict) -> tuple[bool, str]:
        path = self.policy.resolve(args["path"])
        book = _book_plan(args)
        openpyxl = _import_openpyxl()
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                wb = openpyxl.load_workbook(path)
            kinds = chart_kinds(path)
        except Exception as e:  # openpyxl and zipfile raise many types for broken files
            return False, f"the file does not open as an Excel workbook: {e}"
        problems = []
        try:
            want_names = [s.name for s in book.sheets]
            if wb.sheetnames != want_names:
                problems.append(f"sheets {wb.sheetnames} (expected {want_names})")
            else:
                for s in book.sheets:
                    ws = wb[s.name]
                    got = [i for i, row in enumerate(ws.iter_rows(values_only=True), 1)
                           if any(_filled(v) for v in row)]
                    if got != s.filled_rows():
                        problems.append(f"sheet {s.name}: {len(got)} filled rows (expected {len(s.filled_rows())})")
                    header = [ws.cell(row=1, column=c).value for c in range(1, len(s.rows[0]) + 1)]
                    if header != s.rows[0]:
                        problems.append(f"sheet {s.name}: the header row differs")
                    for (r, c), formula in s.formulas.items():
                        cell = ws.cell(row=r, column=c)
                        if cell.data_type != "f" or cell.value != formula:
                            problems.append(f"sheet {s.name}: no formula in {_col_letter(c)}{r}")
                for s, col, fmt in book.formats:
                    if s.last_row >= 2 and wb[s.name].cell(row=2, column=col).number_format != fmt:
                        problems.append(f"sheet {s.name}: column {_col_letter(col)} lacks format {fmt!r}")
        finally:
            wb.close()
        want_kinds = sorted(c.kind for c in book.charts)
        if kinds != want_kinds:
            problems.append(f"charts {kinds or 'none'} (expected {want_kinds or 'none'})")
        if problems:
            return False, "the workbook opens but " + "; ".join(problems)
        rows = sum(len(s.filled_rows()) for s in book.sheets)
        return True, (f"the workbook opens with {_n(len(book.sheets), 'sheet')}, {_n(rows, 'filled row')} and "
                      f"{_n(len(book.charts), 'chart')} as expected")

    # ---- T0 ----

    def analyze_table(self, args: dict) -> str:
        path = self.policy.resolve(args["path"])
        if not path.is_file():
            return f"No file {args['path']}."
        suffix = path.suffix.lower()
        if suffix not in (".csv", ".xlsx"):
            raise OfficeError("analyze_table reads .csv and .xlsx files only")
        size = path.stat().st_size
        if size > MAX_TABLE_BYTES:
            raise OfficeError(f"{args['path']} is too large to analyze ({size} bytes, limit {MAX_TABLE_BYTES})")
        rows, kind = _read_csv(path) if suffix == ".csv" else _read_xlsx(path, args.get("sheet"))
        return analyze_rows(rows, args.get("group_by") or None, args.get("sum_column") or None,
                            f"{self.policy.show(path)} ({kind})")
