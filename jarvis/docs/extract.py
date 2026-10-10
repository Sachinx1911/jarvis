"""Text extraction per page for the document library.

Supported: .md/.markdown/.txt (no pages), .pdf (pypdf, one entry per page), .docx
(python-docx; pages only where the document has explicit page breaks, because Word
decides real page layout at render time). Scanned PDFs have no text layer; OCR comes
in Phase 8, so such pages come back empty and the file is reported as "no text".

pypdf and python-docx are imported lazily so the rest of Jarvis works without them.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED = {".md", ".markdown", ".txt", ".pdf", ".docx"}


class ExtractError(RuntimeError):
    pass


@dataclass(frozen=True)
class Block:
    """A run of text with its location: page (None = no page numbers) and nearest heading."""

    text: str
    page: int | None = None
    section: str = ""


_MD_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$")


def extract(path: Path) -> list[Block]:
    suffix = path.suffix.lower()
    if suffix in (".md", ".markdown", ".txt"):
        return _text(path.read_text(encoding="utf-8", errors="replace"), markdown=suffix != ".txt")
    if suffix == ".pdf":
        return _pdf(path)
    if suffix == ".docx":
        return _docx(path)
    raise ExtractError(f"unsupported file type {suffix}")


def _text(text: str, markdown: bool) -> list[Block]:
    """Split on markdown headings so every chunk can cite its section."""
    if not markdown:
        return [Block(text)] if text.strip() else []
    blocks: list[Block] = []
    section, buf = "", []
    for line in text.splitlines():
        m = _MD_HEADING.match(line)
        if m:
            if "".join(buf).strip():
                blocks.append(Block("\n".join(buf).strip(), None, section))
            section, buf = m.group(1).strip(), [line]
        else:
            buf.append(line)
    if "".join(buf).strip():
        blocks.append(Block("\n".join(buf).strip(), None, section))
    return blocks


def _pdf(path: Path) -> list[Block]:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as e:  # pragma: no cover - dependency missing
        raise ExtractError("pypdf is not installed") from e
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ExtractError("PDF is password protected")
        blocks = []
        for number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                blocks.append(Block(text, number))
        return blocks
    except PdfReadError as e:
        raise ExtractError(f"cannot read PDF: {e}") from e


def _docx(path: Path) -> list[Block]:
    try:
        import docx  # python-docx
        from docx.oxml.ns import qn
    except ImportError as e:  # pragma: no cover - dependency missing
        raise ExtractError("python-docx is not installed") from e
    try:
        document = docx.Document(str(path))
    except Exception as e:  # python-docx raises several types for broken files
        raise ExtractError(f"cannot read DOCX: {e}") from e

    has_breaks = any(
        br.get(qn("w:type")) == "page" for br in document.element.body.iter(qn("w:br"))
    )
    page = 1
    section = ""
    blocks: list[Block] = []
    buf: list[str] = []

    def flush() -> None:
        if "".join(buf).strip():
            blocks.append(Block("\n".join(buf).strip(), page if has_breaks else None, section))
        buf.clear()

    for para in document.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = para.text
        if style.lower().startswith("heading") and text.strip():
            flush()
            section = text.strip()
        if text.strip():
            buf.append(text)
        breaks = sum(1 for br in para._p.iter(qn("w:br")) if br.get(qn("w:type")) == "page")
        if breaks:
            flush()
            page += breaks
    flush()
    for table in document.tables:  # table text is searchable too; cited without a page
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        text = "\n".join(r for r in rows if r.strip(" |"))
        if text.strip():
            blocks.append(Block(text, None, "table"))
    return blocks
