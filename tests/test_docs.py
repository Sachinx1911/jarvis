"""Phase 5 document library: chunking with pages, citations, P2 withholding, re-index on change only."""

import os
import shutil
from pathlib import Path

import pytest

from jarvis.control.audit import AuditLog
from jarvis.db import Store
from jarvis.docs.chunk import chunk_blocks
from jarvis.docs.extract import Block
from jarvis.docs.library import DocError, DocLibrary
from jarvis.docs.tools import doc_context, make_doc_search
from jarvis.providers.base import ToolCall
from jarvis.tools.registry import ToolRegistry

FIXTURES = Path(__file__).parent / "fixtures" / "docs"


def make_pdf(path: Path, pages: list[str]) -> None:
    """A minimal valid PDF with one text page per entry (no PDF library needed to write it)."""
    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objs = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        lines = text.split("\n")
        stream = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({esc(l)}) '" for l in lines) + " ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
                    f"/Contents {5 + 2 * i} 0 R >>".encode())
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    out = b"%PDF-1.4\n"
    offsets = []
    for n, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)


def make_docx(path: Path) -> None:
    import docx
    from docx.enum.text import WD_BREAK

    d = docx.Document()
    d.add_heading("Orchard lease", level=1)
    d.add_paragraph("The lease covers twelve mango trees on the demo plot.")
    d.add_paragraph("End of page one.").runs[0].add_break(WD_BREAK.PAGE)
    d.add_heading("Payment terms", level=2)
    d.add_paragraph("Rent of 900 rupees is due on the fifth of every month.")
    d.save(str(path))


@pytest.fixture
def lib(tmp_path):
    store = Store(tmp_path / "data" / "jarvis.db")
    library = DocLibrary(tmp_path, store)
    docs = tmp_path / "projects" / "garden" / "docs"
    docs.mkdir(parents=True)
    shutil.copy(FIXTURES / "garden-plan.md", docs)
    shutil.copy(FIXTURES / "notes.txt", docs)
    make_pdf(docs / "survey.pdf", ["Soil survey summary\nThe soil is loamy.",
                                   "Page two findings\nNitrogen is low in the north bed.",
                                   "Page three\nRecommendation: add compost."])
    make_docx(docs / "lease.docx")
    yield library, docs, store
    store.close()


def test_chunks_stay_inside_a_page_and_respect_size():
    long = " ".join(f"Sentence number {i} talks about seeds." for i in range(200))
    chunks = chunk_blocks([Block(long, 1, "A"), Block("Short page two.", 2, "B")], size=300, overlap=50)
    assert all(len(c.text) <= 300 + 60 for c in chunks)
    assert {c.page for c in chunks} == {1, 2}
    assert chunks[-1].page == 2 and chunks[-1].text == "Short page two."
    assert [c.ordinal for c in chunks] == list(range(len(chunks)))
    # overlap: the start of a chunk repeats words from the end of the previous one
    assert chunks[1].text.split()[0] in chunks[0].text


def test_index_and_cite_pdf_page_docx_page_and_markdown_section(lib):
    library, docs, _ = lib
    report = library.index()
    assert sorted(report.added) == ["garden/garden-plan.md", "garden/lease.docx", "garden/notes.txt", "garden/survey.pdf"]
    assert not report.failed

    hit = library.search("nitrogen")[0]
    assert (hit.file, hit.page) == ("survey.pdf", 2)
    assert hit.citation() == "survey.pdf, p. 2"

    hit = library.search("rent rupees")[0]
    assert (hit.file, hit.page, hit.section) == ("lease.docx", 2, "Payment terms")

    hit = library.search("seed budget")[0]
    assert hit.file == "garden-plan.md" and hit.section == "Budget"
    assert hit.citation() == 'garden-plan.md, section "Budget"'
    assert hit.doc_date  # file date kept with every chunk


def test_trigram_search_finds_misspelling_and_project_filter(lib):
    library, _, _ = lib
    library.index()
    assert library.search("irigation")[0].file == "garden-plan.md"  # misspelled
    assert library.search("nitrogen", project="other") == []
    with pytest.raises(DocError):
        library.docs_dir("../escape")


def test_reindex_only_changed_files(lib):
    library, docs, store = lib
    first = library.index()
    assert first.chunks > 0

    again = library.index()
    assert len(again.unchanged) == 4 and again.chunks == 0 and not again.added and not again.updated

    # same content, new mtime: hashed, still unchanged
    os.utime(docs / "notes.txt", (1_800_000_000, 1_800_000_000))
    assert "garden/notes.txt" in library.index().unchanged

    (docs / "notes.txt").write_text("The pump warranty was extended to 2030.", encoding="utf-8")
    changed = library.index()
    assert changed.updated == ["garden/notes.txt"] and len(changed.unchanged) == 3
    assert library.search("warranty")[0].text.startswith("The pump warranty was extended")
    assert "March 2028" not in " ".join(h.text for h in library.search("warranty"))

    (docs / "lease.docx").unlink()
    gone = library.index()
    assert gone.removed == ["garden/lease.docx"]
    assert library.search("mango") == []
    assert store.conn.execute("SELECT COUNT(*) FROM doc_chunks c LEFT JOIN doc_files f ON f.id=c.file_id "
                              "WHERE f.id IS NULL").fetchone()[0] == 0


def test_private_by_default_and_withheld_from_cloud_route(lib, tmp_path):
    library, _, store = lib
    library.index()
    cloud = make_doc_search(library, "P1")
    out = cloud.handler({"query": "nitrogen"})
    assert "Nitrogen" not in out
    assert "private chunk" in out and "withheld" in out

    local = make_doc_search(library, "P2")  # a local model route may see private documents
    assert "[survey.pdf, p. 2]" in local.handler({"query": "nitrogen"})

    library.mark("garden", "survey.pdf", "P1")
    out = cloud.handler({"query": "nitrogen"})
    assert "[survey.pdf, p. 2]" in out and "Nitrogen is low" in out
    assert (tmp_path / "projects" / "garden" / "docs_privacy.yaml").read_text(encoding="utf-8").count("survey.pdf: P1") == 1


def test_privacy_file_edits_apply_on_reindex(lib, tmp_path):
    library, _, _ = lib
    library.index()
    (tmp_path / "projects" / "garden" / "docs_privacy.yaml").write_text("files:\n  notes.txt: P1\n", encoding="utf-8")
    library.index()
    assert {d.file: d.privacy for d in library.documents()}["notes.txt"] == "P1"
    assert {d.file: d.privacy for d in library.documents()}["survey.pdf"] == "P2"


def test_mark_rejects_files_outside_docs_and_bad_classes(lib, tmp_path):
    library, _, _ = lib
    (tmp_path / "secret.txt").write_text("x", encoding="utf-8")
    with pytest.raises(DocError):
        library.mark("garden", "../../../secret.txt", "P1")
    with pytest.raises(DocError):
        library.mark("garden", "notes.txt", "P3")
    with pytest.raises(DocError):
        library.mark("garden", "missing.pdf", "P1")


def test_p3_secrets_are_redacted_before_indexing(lib):
    library, docs, store = lib
    (docs / "wifi.md").write_text("# Router\nThe router password is Zq7-demo-pass and the SSID is DemoNet.",
                                  encoding="utf-8")
    library.index()
    library.mark("garden", "wifi.md", "P1")
    out = make_doc_search(library, "P1").handler({"query": "router SSID"})
    assert "Zq7-demo-pass" not in out and "[WITHHELD]" in out
    rows = store.conn.execute("SELECT text FROM doc_chunks").fetchall()
    assert not any("Zq7-demo-pass" in r[0] for r in rows)


def test_hidden_lock_and_unsupported_files_are_skipped_and_empty_pdf_reported(lib):
    library, docs, _ = lib
    (docs / "~$lease.docx").write_bytes(b"lock")
    (docs / ".hidden.md").write_text("hidden", encoding="utf-8")
    (docs / "photo.jpg").write_bytes(b"\xff\xd8")
    make_pdf(docs / "scan.pdf", [""])
    report = library.index()
    names = report.added + report.unchanged
    assert not any(n.endswith(("~$lease.docx", ".hidden.md", "photo.jpg")) for n in names)
    info = {d.file: d for d in library.documents()}
    assert info["scan.pdf"].status == "no_text" and "OCR" in info["scan.pdf"].error


def test_marathi_devanagari_document(lib):
    library, docs, _ = lib
    (docs / "marathi.md").write_text("# बजेट\nबियाण्यांचे बजेट ४५०० रुपये आहे.", encoding="utf-8")
    library.index()
    assert library.search("बजेट")[0].file == "marathi.md"


def test_doc_search_through_registry_runs_at_t0(lib, tmp_path):
    library, _, store = lib
    library.index()
    library.mark("garden", "survey.pdf", "P1")
    reg = ToolRegistry(store, AuditLog(tmp_path / "data" / "audit.jsonl"))  # no approver: T0 must still run
    reg.register(make_doc_search(library, "P1"))
    result = reg.execute(ToolCall("c1", "doc_search", {"query": "compost", "project": "garden"}))
    assert not result.is_error and "[survey.pdf, p. 3]" in result.content


def test_doc_context_hook_counts_withheld(lib):
    library, _, _ = lib
    library.index()
    text, withheld = doc_context(library, "nitrogen in the north bed", "garden", "P1")
    assert text == "" and withheld >= 1
    library.mark("garden", "survey.pdf", "P1")
    text, withheld = doc_context(library, "nitrogen in the north bed", "garden", "P1")
    assert "## From project documents" in text and "[survey.pdf, p. 2]" in text
