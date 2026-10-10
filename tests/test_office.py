"""Office tools: create_docx, create_xlsx, analyze_table (jarvis/tools/office.py).

All data here is made up. Everything runs offline.
"""

import json
import math
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import docx
import openpyxl
import pytest

from jarvis.agents.extra import OPERATOR
from jarvis.agents.team import CHECKER, RESEARCHER, WRITER, AgentEnv, WorkOrder, run_agent
from jarvis.config import load_config
from jarvis.control.approvals import ApprovalGate
from jarvis.control.audit import AuditLog
from jarvis.control.stop import StopSwitch
from jarvis.core.chat import ChatSession
from jarvis.db import Store
from jarvis.memory.store import MemoryStore
from jarvis.providers.base import ToolCall
from jarvis.providers.fake import FakeProvider
from jarvis.routing.policy import RouteDecision
from jarvis.tools import office
from jarvis.tools.catalog import ToolDeps, build_registry
from jarvis.tools.files import PathBlocked, PathDenied, PathPolicy
from jarvis.tools.office import OfficeError, OfficeTools, chart_kinds, check_formula, parse_number
from jarvis.tools.registry import ToolRegistry, output_taints
from jarvis.tools.system import SystemTools
from jarvis.tools.web import WebClient

MAIN = RouteDecision("main_chat", "fake", "fake-1", "medium", "test")


@pytest.fixture
def ws(tmp_path):
    root = tmp_path / "workspace"
    root.mkdir()
    return root


@pytest.fixture
def tools(ws):
    return OfficeTools(PathPolicy([ws]))


def _doc_args(path="out/report.docx", **extra):
    args = {"path": path, "title": "Demo laptop comparison", "blocks": [
        {"type": "heading", "text": "Summary", "level": 1},
        {"type": "paragraph", "text": "Three demo laptops.\nPrices are made up."},
        {"type": "bullets", "items": ["Light", "Long battery", "  "]},
        {"type": "heading", "text": "Steps", "level": 2},
        {"type": "numbered", "items": ["Compare", "Choose"]},
        {"type": "heading", "text": "Next", "level": 3},
        {"type": "numbered", "items": ["Order", "Check", "Return"]},
        {"type": "table", "header": ["Model", "Price"], "rows": [["Demo A", "₹70,000"], ["Demo B"]]},
    ]}
    args.update(extra)
    return args


def _xlsx_args(path="sales.xlsx", **extra):
    args = {
        "path": path,
        "sheets": [
            {"name": "Sales", "rows": [["Month", "Amount", "Units"], ["Jan", 1200.5, 3], ["Feb", 900, 2],
                                       ["Mar", 1500, None], ["=HYPERLINK(\"https://attacker.test\")", 10, 1]]},
            {"name": "Notes", "rows": [["Note"], ["made-up data"]]},
        ],
        "formats": [{"column": "Amount", "format": "#,##0.00"}],
        "formulas": [{"cell": "B7", "formula": "=sum(b2:b5)"}, {"sheet": "Notes", "cell": "B2", "formula": "Sales!B7*2"}],
        "charts": [{"type": "bar", "data": "B1:C5", "categories": "A2:A5", "title": "Sales"},
                   {"type": "pie", "data": "B1:B4", "categories": "A2:A4"},
                   {"type": "line", "sheet": "Sales", "data": "B1:B5", "anchor": "H20"}],
    }
    args.update(extra)
    return args


# ---- create_docx ----

def test_create_docx_writes_structure_and_verifies(tools, ws):
    args = _doc_args()
    msg = tools.create_docx(args)
    assert msg.startswith("Created out/report.docx:") and "3 headings" in msg and "1 table" in msg
    assert tools.verify_docx(args) == (True, "the document opens with 3 headings, 1 table and 7 list items as expected")
    doc = docx.Document(str(ws / "out" / "report.docx"))
    styles = [(p.style.name, p.text) for p in doc.paragraphs]
    assert styles[0] == ("Title", "Demo laptop comparison")
    assert ("Heading 1", "Summary") in styles and ("Heading 3", "Next") in styles
    assert [t for s, t in styles if s == "List Bullet"] == ["Light", "Long battery"]  # blank item dropped
    assert "\n" in next(t for s, t in styles if s == "Normal")  # line break kept inside the paragraph
    # each numbered list restarts at 1: its own numbering instance
    num_ids = {p._p.pPr.numPr.numId.val for p in doc.paragraphs if p.style.name == "List Number"}
    assert len(num_ids) == 2
    table = doc.tables[0]
    assert [[c.text for c in r.cells] for r in table.rows] == [["Model", "Price"], ["Demo A", "₹70,000"], ["Demo B", ""]]
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold
    assert doc.core_properties.title == "Demo laptop comparison"


def test_create_docx_never_overwrites(tools, ws):
    (ws / "report.docx").write_bytes(b"keep me")
    with pytest.raises(OfficeError, match="already exists"):
        tools.create_docx(_doc_args("report.docx"))
    assert (ws / "report.docx").read_bytes() == b"keep me"


def test_exclusive_create_refuses_a_file_that_appeared_meanwhile(ws):
    target = ws / "late.docx"
    target.write_bytes(b"other process")
    with pytest.raises(OfficeError, match="already exists"):
        office._write_new(target, b"new bytes")
    assert target.read_bytes() == b"other process"


@pytest.mark.parametrize("args,error", [
    ({"path": "a.txt", "blocks": [{"type": "paragraph", "text": "x"}]}, "must end with .docx"),
    ({"path": "a.docx", "blocks": []}, "nothing to write"),
    ({"path": "a.docx", "blocks": [{"type": "heading", "text": "x", "level": 4}]}, "level must be 1, 2 or 3"),
    ({"path": "a.docx", "blocks": [{"type": "paragraph", "text": "  "}]}, "has no text"),
    ({"path": "a.docx", "blocks": [{"type": "bullets", "items": []}]}, "has no items"),
    ({"path": "a.docx", "blocks": [{"type": "table", "rows": [["x"]]}]}, "needs a header row"),
    ({"path": "a.docx", "blocks": [{"type": "table", "header": ["a"], "rows": [["x", "y"]]}]}, "has 2 cells"),
    ({"path": "a.docx", "blocks": [{"type": "heading", "text": "x" * 301}]}, "too long"),
])
def test_create_docx_rejects_bad_input(tools, ws, args, error):
    with pytest.raises(OfficeError, match=error):
        tools.create_docx(args)
    assert not any(ws.rglob("*.docx"))


def test_create_docx_size_caps(tools, monkeypatch):
    monkeypatch.setattr(office, "MAX_DOC_BLOCKS", 2)
    with pytest.raises(OfficeError, match="too many blocks: 3"):
        tools.create_docx({"path": "a.docx", "blocks": [{"type": "paragraph", "text": "x"}] * 3})
    monkeypatch.setattr(office, "MAX_DOC_CHARS", 50)
    with pytest.raises(OfficeError, match="too long"):
        tools.create_docx({"path": "b.docx", "blocks": [{"type": "paragraph", "text": "y" * 60}]})
    monkeypatch.setattr(office, "MAX_TABLE_CELLS", 5)
    with pytest.raises(OfficeError, match="too many table cells"):
        tools.create_docx({"path": "c.docx", "blocks": [{"type": "table", "header": ["a", "b"], "rows": [["1", "2"]] * 2}]})


def test_verify_docx_fails_when_the_file_does_not_match(tools, ws):
    tools.create_docx(_doc_args("v.docx"))
    more = _doc_args("v.docx")
    more["blocks"] = more["blocks"] + [{"type": "heading", "text": "Extra"}, {"type": "table", "header": ["x"]}]
    ok, how = tools.verify_docx(more)
    assert not ok and "headings 3 (expected 4)" in how and "tables" in how
    (ws / "broken.docx").write_bytes(b"not a zip")
    ok, how = tools.verify_docx(_doc_args("broken.docx"))
    assert not ok and "does not open" in how


def test_control_characters_are_removed_not_fatal(tools, ws):
    args = {"path": "c.docx", "blocks": [{"type": "paragraph", "text": "bell\x07 and nul\x00 gone"}]}
    tools.create_docx(args)
    assert docx.Document(str(ws / "c.docx")).paragraphs[0].text == "bell and nul gone"


# ---- create_xlsx ----

def test_create_xlsx_sheets_formats_formulas_charts(tools, ws):
    args = _xlsx_args()
    msg = tools.create_xlsx(args)
    assert msg == ("Created sales.xlsx: sheet Sales (6 rows), sheet Notes (2 rows); 2 formulas; "
                   "3 charts (1 bar, 1 line, 1 pie).")
    ok, how = tools.verify_xlsx(args)
    assert ok, how
    path = ws / "sales.xlsx"
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Sales", "Notes"]
    s = wb["Sales"]
    assert [c.value for c in s[1]] == ["Month", "Amount", "Units"] and s["A1"].font.bold
    assert s["B2"].value == 1200.5 and s["B2"].number_format == "#,##0.00"
    assert s["B7"].value == "=SUM(B2:B5)" and s["B7"].data_type == "f"  # explicit formula, in capitals
    assert wb["Notes"]["B2"].value == "=Sales!B7*2"
    # text that looks like a formula stays text: web-derived data never becomes a live formula
    assert s["A5"].value == '=HYPERLINK("https://attacker.test")' and s["A5"].data_type == "s"
    assert s.freeze_panes == "A2"
    assert chart_kinds(path) == ["bar", "line", "pie"]
    with zipfile.ZipFile(path) as z:  # native charts, no images
        assert not [n for n in z.namelist() if n.startswith("xl/media/")]


def test_create_xlsx_never_overwrites(tools, ws):
    (ws / "sales.xlsx").write_bytes(b"keep")
    with pytest.raises(OfficeError, match="already exists"):
        tools.create_xlsx(_xlsx_args())
    assert (ws / "sales.xlsx").read_bytes() == b"keep"


def _with(**changes):
    args = _xlsx_args()
    args.update(changes)
    return args


@pytest.mark.parametrize("args,error", [
    (_with(path="sales.csv"), "must end with .xlsx"),
    (_with(formulas=[{"cell": "D2", "formula": '=HYPERLINK("https://attacker.test","x")'}]), "HYPERLINK is not allowed"),
    (_with(formulas=[{"cell": "D2", "formula": '=WEBSERVICE("https://attacker.test")'}]), "WEBSERVICE is not allowed"),
    (_with(formulas=[{"cell": "D2", "formula": "=INDIRECT(A1)"}]), "INDIRECT is not allowed"),
    (_with(formulas=[{"cell": "D2", "formula": "=cmd|' /C calc'!A0"}]), "is not allowed near"),
    (_with(formulas=[{"cell": "D2", "formula": "=[1]Sheet1!A1"}]), "is not allowed near"),
    (_with(formulas=[{"cell": "D2", "formula": "=Other!A1"}]), "unknown sheet"),
    (_with(formulas=[{"cell": "D2", "formula": "=SUM(B2:B5"}]), "not well formed"),
    (_with(formulas=[{"cell": "D2", "formula": "=B2 B3"}]), "not well formed"),
    (_with(formulas=[{"cell": "B2", "formula": "=1+1"}]), "already holds a value"),
    (_with(formulas=[{"cell": "nowhere", "formula": "=1+1"}]), "is not a cell"),
    (_with(sheets=[{"name": "Bad/Name", "rows": [["a"]]}], formats=[], formulas=[], charts=[]), "not allowed in Excel"),
    (_with(sheets=[{"name": "A", "rows": [["a"]]}, {"name": "a", "rows": [["b"]]}], formats=[], formulas=[], charts=[]),
     "two sheets are named"),
    (_with(sheets=[{"name": "A", "rows": [[None, ""], [1, 2]]}], formats=[], formulas=[], charts=[]), r"header\) is empty"),
    (_with(sheets=[{"name": "A", "rows": [["a"], [math.nan]]}], formats=[], formulas=[], charts=[]), "finite"),
    (_with(sheets=[{"name": "A", "rows": [["a"], [{"x": 1}]]}], formats=[], formulas=[], charts=[]), "cells must be"),
    (_with(sheets=[], formats=[], formulas=[], charts=[]), "at least one sheet"),
    (_with(charts=[{"type": "pie", "data": "B1:C5"}]), "one column of values"),
    (_with(charts=[{"type": "bar", "data": "B1:Z5"}]), "outside the data"),
    (_with(charts=[{"type": "bar", "data": "B1:B1"}]), "at least one row of values"),
    (_with(charts=[{"type": "bar", "data": "B1:B5", "sheet": "Missing"}]), "no sheet named"),
    (_with(formats=[{"column": "Price", "format": "0.00"}]), "no column 'Price'"),
    (_with(formats=[{"column": "B", "format": '"Rs'}]), "not valid"),
])
def test_create_xlsx_rejects_bad_input(tools, ws, args, error):
    with pytest.raises(OfficeError, match=error):
        tools.create_xlsx(args)
    assert not any(ws.rglob("*.xlsx"))


def test_create_xlsx_size_caps(tools, monkeypatch):
    monkeypatch.setattr(office, "MAX_SHEET_ROWS", 3)
    with pytest.raises(OfficeError, match="too many rows in sheet 'Sales': 5"):
        tools.create_xlsx(_xlsx_args())
    monkeypatch.setattr(office, "MAX_SHEET_ROWS", 10_000)
    monkeypatch.setattr(office, "MAX_CELLS", 10)
    with pytest.raises(OfficeError, match="too many cells in the workbook"):
        tools.create_xlsx(_xlsx_args())
    monkeypatch.setattr(office, "MAX_CELLS", 200_000)
    monkeypatch.setattr(office, "MAX_CHARTS", 2)
    with pytest.raises(OfficeError, match="too many charts: 3"):
        tools.create_xlsx(_xlsx_args())


def test_verify_xlsx_fails_when_the_file_does_not_match(tools, ws):
    tools.create_xlsx(_xlsx_args())
    extra_chart = _xlsx_args(charts=_xlsx_args()["charts"] + [{"type": "bar", "data": "C1:C5"}])
    ok, how = tools.verify_xlsx(extra_chart)
    assert not ok and "charts" in how
    extra_row = _xlsx_args()
    extra_row["sheets"][1]["rows"].append(["one more"])
    ok, how = tools.verify_xlsx(extra_row)
    assert not ok and "sheet Notes: 2 filled rows (expected 3)" in how
    (ws / "broken.xlsx").write_bytes(b"not a zip")
    ok, how = tools.verify_xlsx(_xlsx_args("broken.xlsx"))
    assert not ok and "does not open" in how


@pytest.mark.parametrize("formula,expected", [
    ("SUM(B2:B5)", "=SUM(B2:B5)"),
    ("=average(Sales!b2:b5)", "=AVERAGE(Sales!B2:B5)"),
    ('=IF(B2>=100,"big","small")', '=IF(B2>=100,"big","small")'),
    ("=-B2*2+ROUND(C3/4,2)", "=-B2*2+ROUND(C3/4,2)"),
    ("=$B$2*10%", "=$B$2*10%"),
])
def test_check_formula_allows_simple_formulas(formula, expected):
    assert check_formula(formula, {"sales"}) == expected


# ---- analyze_table ----

@pytest.mark.parametrize("text,value", [
    ("1,23,456.50", Decimal("123456.50")),
    ("12,34,56,789", Decimal("123456789")),
    ("123,456.50", Decimal("123456.50")),
    ("₹1,234", Decimal("1234")),
    ("₹ 1,23,456/-", Decimal("123456")),
    ("Rs. 500", Decimal("500")),
    ("rs 75.25", Decimal("75.25")),
    ("INR 1,000", Decimal("1000")),
    ("(1,200)", Decimal("-1200")),
    ("-₹50", Decimal("-50")),
    ("₹-50", Decimal("-50")),
    ("−7", Decimal("-7")),
    (".5", Decimal("0.5")),
    ("42", Decimal("42")),
    ("12,3", None), ("1,2,3", None), ("1.2.3", None), ("+-5", None), ("abc", None), ("", None), ("₹", None),
    ("1,23,4567", None), ("10%", None),
])
def test_parse_number_understands_indian_formats(text, value):
    assert parse_number(text) == value


CSV = (
    "Region,Amount,Note\n"
    'North,"1,23,456.50",first\n'
    "South,₹1234,second\n"
    "North,Rs. 500,\n"
    ",,\n"
    "East,(200),third\n"
    'West,"₹ 2,00,000/-",fourth\n'
)


def test_analyze_csv_sums_exactly_with_group_by(tools, ws):
    (ws / "sales.csv").write_text(CSV, encoding="utf-8")
    out = tools.analyze_table({"path": "sales.csv", "group_by": "region", "sum_column": "AMOUNT"})
    lines = out.splitlines()
    assert lines[0] == "Table: sales.csv (CSV, comma-separated)"
    assert "Rows: 5 data rows (header excluded), 3 columns. 1 blank row skipped." in lines
    assert "- Amount: number, 5 values; sum 324990.5, mean 64998.1, min -200, max 200000" in lines
    assert "- Region: text, 5 values, 4 distinct" in lines
    assert "- Note: text, 4 values, 4 distinct, 1 blank" in lines
    i = lines.index("Group by Region (sum of Amount):")
    assert lines[i + 1:i + 5] == ["- West: sum 200000, 1 row", "- North: sum 123956.5, 2 rows",
                                  "- South: sum 1234, 1 row", "- East: sum -200, 1 row"]
    assert "Total: sum 324990.5 over 5 rows in 4 groups" in lines
    assert lines[-1] == "All numbers above were computed in code from the file."


def test_analyze_counts_groups_flags_total_rows_and_mixed_columns(tools, ws):
    (ws / "t.csv").write_text("City;Value\nPune;10\nPune;n/a\nNashik;x\nMumbai;5\nTotal;15\n", encoding="utf-8")
    out = tools.analyze_table({"path": "t.csv", "group_by": "City"})
    assert "(CSV, semicolon-separated)" in out
    assert "- Value: number, 3 values, 1 non-numeric skipped; sum 30, mean 10, min 5, max 15, 1 blank" in out
    assert "Group by City (row count):\n- Pune: 2 rows\n" in out and "Total: 5 rows in 4 groups" in out
    assert "Note: row 6 looks like a total row ('Total'); it is included in the sums above." in out


def test_analyze_group_keys_are_never_merged_by_truncation(tools, ws):
    long_a, long_b = "Demo customer " + "x" * 60 + "A", "Demo customer " + "x" * 60 + "B"
    (ws / "k.csv").write_text(f"Who,Amt\n{long_a},1\n{long_b},2\n", encoding="utf-8")
    out = tools.analyze_table({"path": "k.csv", "group_by": "Who", "sum_column": "Amt"})
    assert "in 2 groups" in out


def test_analyze_xlsx_with_sheet_choice(tools, ws):
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    wb.active.append(["nothing here"])
    data = wb.create_sheet("Data")
    for row in (["Item", "Qty", "Price"], ["Pen", 2, 10.5], ["Book", 1, "₹1,250"], ["Pen", 3, "1,00,000.25"]):
        data.append(row)
    data.append(["Day", 1, datetime(2026, 10, 10)])
    wb.save(ws / "stock.xlsx")
    out = tools.analyze_table({"path": "stock.xlsx", "sheet": "data", "group_by": "Item", "sum_column": "Price"})
    assert "Table: stock.xlsx (XLSX, sheet Data (2 of 2: Cover, Data))" in out
    assert "- Qty: number, 4 values; sum 7, mean 1.75, min 1, max 3" in out
    assert "- Price: number, 3 values, 1 non-numeric skipped; sum 101260.75" in out
    assert "- Pen: sum 100010.75, 2 rows" in out and "- Book: sum 1250, 1 row" in out
    with pytest.raises(OfficeError, match="no sheet named 'Missing'"):
        tools.analyze_table({"path": "stock.xlsx", "sheet": "Missing"})


def test_analyze_errors_are_clear(tools, ws, monkeypatch):
    (ws / "t.csv").write_text("A,B\nx,1\ny,2\n", encoding="utf-8")
    (ws / "notes.txt").write_text("hello", encoding="utf-8")
    assert tools.analyze_table({"path": "missing.csv"}) == "No file missing.csv."
    with pytest.raises(OfficeError, match=r"\.csv and \.xlsx"):
        tools.analyze_table({"path": "notes.txt"})
    with pytest.raises(OfficeError, match=r"no group_by column 'C' \(columns: A, B\)"):
        tools.analyze_table({"path": "t.csv", "group_by": "C"})
    with pytest.raises(OfficeError, match="not numeric"):
        tools.analyze_table({"path": "t.csv", "group_by": "B", "sum_column": "A"})
    with pytest.raises(OfficeError, match="sum_column needs group_by"):
        tools.analyze_table({"path": "t.csv", "sum_column": "B"})
    (ws / "bad.xlsx").write_bytes(b"not a zip")
    with pytest.raises(OfficeError, match="cannot read bad.xlsx"):
        tools.analyze_table({"path": "bad.xlsx"})
    monkeypatch.setattr(office, "MAX_ANALYZE_ROWS", 1)
    with pytest.raises(OfficeError, match="too many rows to analyze"):
        tools.analyze_table({"path": "t.csv"})
    monkeypatch.setattr(office, "MAX_TABLE_BYTES", 5)
    with pytest.raises(OfficeError, match="too large"):
        tools.analyze_table({"path": "t.csv"})


def test_paths_stay_inside_allowed_folders(tools, ws, tmp_path):
    (tmp_path / "outside.csv").write_text("A\n1\n", encoding="utf-8")
    for call, args in ((tools.create_docx, _doc_args("../escape.docx")), (tools.create_xlsx, _xlsx_args("/tmp/x.xlsx")),
                       (tools.analyze_table, {"path": str(tmp_path / "outside.csv")})):
        with pytest.raises(PathDenied):
            call(args)
    for call, args in ((tools.create_docx, _doc_args(".ssh/notes.docx")), (tools.create_xlsx, _xlsx_args(".env.xlsx")),
                       (tools.analyze_table, {"path": ".aws/costs.csv"})):
        with pytest.raises(PathBlocked):
            call(args)
    assert not (tmp_path / "escape.docx").exists()


# ---- registry: tiers, verify, taint ----

def _registry(home, files_allowed=None):
    store = Store(home / "data" / "jarvis.db")
    audit = AuditLog(home / "data" / "audit.jsonl")
    reg = ToolRegistry(store, audit, ApprovalGate(store, audit))
    tools_config = {"files": {"allowed": files_allowed}} if files_allowed else {}
    build_registry(reg, ToolDeps(home=home, store=store, audit=audit, memory=None, skills=None,
                                 session_id=lambda: "s", max_privacy="P1", tools_config=tools_config, web=WebClient(),
                                 quarantine=None, tavily_key=lambda: None, system=SystemTools()))
    return reg, store


def test_tiers_verify_and_taint_choices(isolated):
    reg, store = _registry(isolated)
    docx_tool, xlsx_tool, analyze = reg.get("create_docx"), reg.get("create_xlsx"), reg.get("analyze_table")
    assert (docx_tool.tier, xlsx_tool.tier, analyze.tier) == ("T1", "T1", "T0")
    assert docx_tool.verify and xlsx_tool.verify and docx_tool.writes and xlsx_tool.writes
    assert docx_tool.precheck and xlsx_tool.precheck and analyze.precheck
    assert not output_taints(docx_tool) and not output_taints(xlsx_tool)  # Jarvis's own confirmation text
    assert analyze.taint_check is not None and analyze.verify is None
    names = {s.name for s in reg.specs()}
    assert {"create_docx", "create_xlsx", "analyze_table"} <= names
    store.close()


def test_analyze_taints_only_for_untrusted_files(isolated):
    reg, store = _registry(isolated)
    root = isolated / "workspace"
    (root / "mine.csv").write_text("A,B\nx,1\n", encoding="utf-8")
    (root / "downloaded.csv").write_text("A,B\ny,2\n", encoding="utf-8")
    reg.provenance.record(root / "downloaded.csv", True, "operator:browser_download")
    res = reg.execute(ToolCall("c1", "analyze_table", {"path": "mine.csv"}))
    assert not res.is_error and not reg.tainted  # the user's own file stays clean
    res = reg.execute(ToolCall("c2", "analyze_table", {"path": "downloaded.csv"}))
    assert not res.is_error and "sum 2" in res.content and reg.tainted
    store.close()


def test_analyze_taints_for_files_in_task_folders(isolated):
    task_dir = isolated / "data" / "tasks" / "T9"
    task_dir.mkdir(parents=True)
    (task_dir / "prices.csv").write_text("Shop,Price\nDemo,\"₹1,000\"\n", encoding="utf-8")
    reg, store = _registry(isolated, [str(task_dir)])
    reg.execute(ToolCall("c1", "analyze_table", {"path": "prices.csv"}))
    assert reg.tainted
    store.close()


def _events(home):
    return [json.loads(l) for l in (home / "data" / "audit.jsonl").read_text(encoding="utf-8").splitlines()]


def _chat(home, script):
    store = Store(home / "data" / "jarvis.db")
    audit = AuditLog(home / "data" / "audit.jsonl")
    s = ChatSession(load_config(), FakeProvider(script=script), MAIN, store, audit,
                    now=lambda: datetime(2026, 10, 10, 19, 30), home=home, memory=MemoryStore(home, store),
                    approver=None, quarantine=None, web=WebClient(transport=lambda *a: (404, b""), check_url=lambda u: None),
                    system=SystemTools(apps={}, runner=lambda a: None, opener=lambda u: True))
    return s, store


def test_chat_creates_office_files_without_asking_and_verifies(isolated):
    home = isolated
    s, store = _chat(home, [{"tool": "create_docx", "arguments": _doc_args("notes/plan.docx")},
                            {"tool": "create_xlsx", "arguments": _xlsx_args("notes/sales.xlsx")},
                            {"tool": "create_docx", "arguments": _doc_args("notes/plan.docx")},  # exists now
                            "Done, sir."])
    assert s.send("make the plan and the sheet") == "Done, sir."
    runs = [e for e in _events(home) if e["event"] == "tool_run"]
    assert [(r["tool"], r["decision"], r["ok"], r["verified"]) for r in runs] == [
        ("create_docx", "auto", True, True), ("create_xlsx", "auto", True, True), ("create_docx", "auto", False, None)]
    assert (home / "workspace" / "notes" / "plan.docx").is_file() and (home / "workspace" / "notes" / "sales.xlsx").is_file()
    result = store.conn.execute("SELECT result FROM tool_runs WHERE tool='create_docx' ORDER BY id DESC").fetchone()[0]
    assert "already exists" in result
    store.close()


def test_chat_refuses_secret_and_outside_paths(isolated):
    home = isolated
    s, store = _chat(home, [{"tool": "create_docx", "arguments": _doc_args(".ssh/x.docx")},
                            {"tool": "create_xlsx", "arguments": _xlsx_args("../outside.xlsx")}, "ok"])
    s.send("write them")
    ev = _events(home)
    runs = [e for e in ev if e["event"] == "tool_run"]
    assert [r["decision"] for r in runs] == ["blocked_t4", "refused"]
    assert sum(e["event"] == "t4_attempt" for e in ev) == 1
    assert not (home / "outside.xlsx").exists()
    store.close()


# ---- the Writer agent ----

def test_writer_allowlist_has_office_tools_and_nobody_else_gains_writes():
    assert {"create_docx", "create_xlsx", "analyze_table"} <= WRITER.tools
    for spec in (RESEARCHER, CHECKER, OPERATOR):
        assert not {"create_docx", "create_xlsx"} & spec.tools


def test_writer_creates_office_files_sandboxed_in_the_task_folder(isolated):
    home = isolated
    task_dir = home / "data" / "tasks" / "T1"
    task_dir.mkdir(parents=True)
    (task_dir / "research-1.csv").write_text("Model,Price\nDemo A,\"₹70,000\"\nDemo B,\"₹65,500\"\n", encoding="utf-8")
    provider = FakeProvider(script=[
        {"tool": "analyze_table", "arguments": {"path": "research-1.csv"}},
        {"tool": "create_docx", "arguments": _doc_args("report.docx")},
        {"tool": "create_xlsx", "arguments": _xlsx_args("report.xlsx")},
        {"tool": "create_docx", "arguments": _doc_args(str(home / "workspace" / "escape.docx"))},  # outside: refused
        "report.docx and report.xlsx created",
    ])
    env = AgentEnv(load_config(), home, home / "data" / "jarvis.db", AuditLog(home / "data" / "audit.jsonl"),
                   StopSwitch(home / "data"), "T1", task_dir, lambda d: provider,
                   lambda r: RouteDecision(r, "fake", "fake-1", "low", "test"), lambda: None, web=WebClient())
    run_agent(WRITER, "writer", WorkOrder("Write the report as Word and Excel files"), env)
    assert (task_dir / "report.docx").is_file() and (task_dir / "report.xlsx").is_file()
    assert not (home / "workspace" / "escape.docx").exists()
    runs = [e for e in _events(home) if e["event"] == "tool_run"]
    assert [(r["tool"], r["decision"], r["ok"]) for r in runs] == [
        ("analyze_table", "auto", True), ("create_docx", "auto_sandboxed", True),
        ("create_xlsx", "auto_sandboxed", True), ("create_docx", "refused", False)]
    assert runs[1]["verified"] is True and runs[2]["verified"] is True
    store = Store(home / "data" / "jarvis.db")
    rows = store.conn.execute("SELECT source, tainted FROM file_provenance ORDER BY id").fetchall()
    store.close()
    assert rows == [("writer:create_docx", 1), ("writer:create_xlsx", 1)]  # what a tainted writer makes stays tainted


def test_receipt_lists_office_files_created_by_a_task(tmp_path):
    from types import SimpleNamespace

    from jarvis.db import Store
    from jarvis.tasks.receipt import build_receipt
    from jarvis.tasks.store import TaskStore

    store = Store(tmp_path / "data" / "jarvis.db")
    try:
        tasks = TaskStore(store, tmp_path / "data" / "tasks")
        t = tasks.create("make the example report", title="Example report")
        for tool, path in (("create_docx", "report.docx"), ("create_xlsx", "table.xlsx"), ("write_file", "notes.md")):
            store.conn.execute("INSERT INTO tool_runs (task_id, tool, tier, arguments, ok) VALUES (?, ?, 'T1', ?, 1)",
                               (t.id, tool, f'{{"path": "{path}"}}'))
        store.conn.commit()
        text = build_receipt(SimpleNamespace(tasks=tasks, store=store), t.id)
    finally:
        store.close()
    assert "created report.docx" in text and "created table.xlsx" in text and "created notes.md" in text
